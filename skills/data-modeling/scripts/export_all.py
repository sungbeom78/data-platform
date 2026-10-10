# model.json → 02_ddl_postgresql.sql / 03_model.dbm / 04_table_spec.xlsx
# 사용: python export_all.py model.json [출력폴더] [--pgmodeler-ver 1.1.0]
#
# v2 확장(모두 선택 키 — 없으면 v1 과 바이트 단위로 같은 결과를 낸다):
#   model.pi_schema        : "pi"  → pi_ 테이블을 별도 스키마에 생성
#   model.meta_mode        : "shared" → meta 사전 테이블을 IF NOT EXISTS 로, 시드는 ON CONFLICT DO NOTHING
#   model.access_policies  : [{schema, table, min_grade, category}] → meta.object_policy 등록(플랫폼 설치 시)
#   table.schema           : 테이블별 스키마(기본 model.schema)
#   table.uniques          : [["col_a","col_b"], ...] 복합 UNIQUE
#   table.checks           : [{"name":"x","expr":"..."}] 테이블 CHECK
#   table.seed_cols        : seed 에 쓸 컬럼 목록(기본 앞 3개)
#   column.crypto          : aes | hmac | argon2  (개인정보 암호화 방식 — COMMENT·정의서에 표시)
import json, re, sys, os, argparse
from xml.sax.saxutils import escape

ap = argparse.ArgumentParser()
ap.add_argument("model"); ap.add_argument("outdir", nargs="?", default=".")
ap.add_argument("--pgmodeler-ver", default="1.1.0")
A = ap.parse_args()
os.makedirs(A.outdir, exist_ok=True); os.chdir(A.outdir) if False else None
M = json.load(open(A.model, encoding="utf-8"))
OUT = lambda f: os.path.join(A.outdir, f)
S, MS = M["schema"], M["meta_schema"]
PIS = M.get("pi_schema")
SHARED = M.get("meta_mode") == "shared"
DOM = {d["name"]: d for d in M["domains"]}
TBL = {t["name"]: t for t in M["tables"]}
CRYPTO_LABEL = {"aes": "AES-128 양방향", "hmac": "HMAC-SHA256 단방향(검색)", "argon2": "argon2id 단방향(인증)"}
def ctype(c): return DOM[c["domain"]]["type"]
def q(s): return "'" + str(s).replace("'", "''") + "'"
def pkcols(t): return [c["name"] for c in t["columns"] if c.get("pk")]
def TS(t):
    """테이블이 속하는 스키마"""
    if t.get("schema"): return t["schema"]
    if PIS and t["name"].startswith("pi_"): return PIS
    return S
def TSN(name): return TS(TBL[name]) if name in TBL else S
SCHEMAS = []
for _s in [S] + [TS(t) for t in M["tables"]]:
    if _s not in SCHEMAS and _s != MS: SCHEMAS.append(_s)
EXTRA = SCHEMAS[1:]
OC = " ON CONFLICT DO NOTHING" if SHARED else ""

# ---------------- 1. DDL ----------------
L = []
w = L.append
w(f"-- {M['project']} v{M['version']} — PostgreSQL {M['db_version']} DDL (generated {M['generated_at']})")
w("-- 생성 순서: 스키마 → meta 사전 → 사전 시드 → 코드 테이블 → 기준/거래/PI 테이블 → FK → 인덱스 → COMMENT")
w("BEGIN;\n")
w(f"CREATE SCHEMA IF NOT EXISTS {S};")
for _s in EXTRA: w(f"CREATE SCHEMA IF NOT EXISTS {_s};")
w(f"CREATE SCHEMA IF NOT EXISTS {MS};\n")

META_BLOCK = f"""-- ===== meta: 용어/단어 사전 =====
CREATE TABLE {MS}.word (
  word_id      bigint GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
  kor_name     varchar(50)  NOT NULL UNIQUE,
  eng_abbr     varchar(30)  NOT NULL UNIQUE,
  eng_full     varchar(60)  NOT NULL,
  description  text,
  is_class     boolean NOT NULL DEFAULT false,
  created_at   timestamptz NOT NULL DEFAULT now()
);
CREATE TABLE {MS}.domain (
  domain_id    bigint GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
  domain_name  varchar(50)  NOT NULL UNIQUE,
  data_type    varchar(50)  NOT NULL,
  check_rule   text,
  example      varchar(100),
  description  text
);
CREATE TABLE {MS}.term (
  term_id      bigint GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
  kor_name     varchar(100) NOT NULL UNIQUE,
  phys_name    varchar(63)  NOT NULL UNIQUE,
  word_path    varchar(200) NOT NULL,
  domain_id    bigint REFERENCES {MS}.domain(domain_id),
  description  text
);
CREATE TABLE {MS}.table_catalog (
  schema_name  varchar(63) NOT NULL,
  table_name   varchar(63) NOT NULL,
  kor_name     varchar(100) NOT NULL,
  table_kind   varchar(20),
  description  text,
  is_pi        boolean NOT NULL DEFAULT false,
  PRIMARY KEY (schema_name, table_name)
);
CREATE TABLE {MS}.column_catalog (
  schema_name  varchar(63) NOT NULL,
  table_name   varchar(63) NOT NULL,
  column_name  varchar(63) NOT NULL,
  ordinal      int NOT NULL,
  term_id      bigint NOT NULL REFERENCES {MS}.term(term_id),
  is_pi        boolean NOT NULL DEFAULT false,
  PRIMARY KEY (schema_name, table_name, column_name),
  FOREIGN KEY (schema_name, table_name) REFERENCES {MS}.table_catalog(schema_name, table_name)
);
COMMENT ON TABLE {MS}.word IS '단어 사전: 이름을 구성하는 최소 단위';
COMMENT ON TABLE {MS}.domain IS '도메인: 자료형과 검증 규칙의 묶음';
COMMENT ON TABLE {MS}.term IS '용어 사전: 단어 조합 = 실제 테이블/컬럼 물리명';
COMMENT ON TABLE {MS}.table_catalog IS '테이블 목록(논리명·구분·개인정보 여부)';
COMMENT ON TABLE {MS}.column_catalog IS '컬럼 목록(용어 참조·개인정보 여부)';
"""
if SHARED:
    META_BLOCK = META_BLOCK.replace("CREATE TABLE ", "CREATE TABLE IF NOT EXISTS ").replace(
        "-- ===== meta: 용어/단어 사전 =====", "-- ===== meta: 용어/단어 사전 (공유 모드: 이미 있으면 재사용) =====")
w(META_BLOCK)

w("-- ----- 사전 시드 -----")
for x in M["words"]:
    w(f"INSERT INTO {MS}.word (kor_name, eng_abbr, eng_full, description, is_class) VALUES ({q(x['kor'])}, {q(x['abbr'])}, {q(x['full'])}, {q(x['desc'])}, {str(x['is_class']).lower()}){OC};")
for d in M["domains"]:
    w(f"INSERT INTO {MS}.domain (domain_name, data_type, check_rule, example, description) VALUES ({q(d['name'])}, {q(d['type'])}, {q(d['check']) if d['check'] else 'NULL'}, {q(d['example'])}, {q(d['desc'])}){OC};")
for t in M["terms"]:
    dom = f"(SELECT domain_id FROM {MS}.domain WHERE domain_name = {q(t['domain'])})" if t["domain"] else "NULL"
    w(f"INSERT INTO {MS}.term (kor_name, phys_name, word_path, domain_id, description) VALUES ({q(t['kor'])}, {q(t['phys'])}, {q(t['word_path'])}, {dom}, {q(t['desc']) if t.get('desc') else 'NULL'}){OC};")
for t in M["tables"]:
    w(f"INSERT INTO {MS}.table_catalog (schema_name, table_name, kor_name, table_kind, description, is_pi) VALUES ({q(TS(t))}, {q(t['name'])}, {q(t['kor'])}, {q(t['kind'])}, {q(t['description'])}, {str(t['is_pi']).lower()}){OC};")
    for i, c in enumerate(t["columns"], 1):
        w(f"INSERT INTO {MS}.column_catalog (schema_name, table_name, column_name, ordinal, term_id, is_pi) VALUES ({q(TS(t))}, {q(t['name'])}, {q(c['name'])}, {i}, (SELECT term_id FROM {MS}.term WHERE phys_name = {q(c['name'])}), {str(bool(c.get('pi'))).lower()}){OC};")
w("")

w(f"-- ===== {S}: 업무 테이블 =====" if not EXTRA else f"-- ===== {', '.join(SCHEMAS)}: 업무 테이블 =====")
for t in M["tables"]:
    w(f"-- {t['kor']} ({t['kind']})")
    lines = []
    for c in t["columns"]:
        s = f"  {c['name']:<14} {ctype(c)}"
        if c.get("identity"): s += " GENERATED ALWAYS AS IDENTITY"
        if not c["nullable"]: s += " NOT NULL"
        if c.get("default") is not None: s += f" DEFAULT {c['default']}"
        lines.append(s)
    lines.append(f"  CONSTRAINT pk_{t['name']} PRIMARY KEY ({', '.join(pkcols(t))})")
    for c in t["columns"]:
        if c.get("unique"):
            lines.append(f"  CONSTRAINT uk_{t['name']}_{c['name']} UNIQUE ({c['name']})")
        chk = DOM[c["domain"]]["check"]
        if chk:
            lines.append(f"  CONSTRAINT ck_{t['name']}_{c['name']} CHECK ({c['name']} {chk})")
    for u in t.get("uniques", []):
        lines.append(f"  CONSTRAINT uk_{t['name']}_{'_'.join(u)} UNIQUE ({', '.join(u)})")
    for ck in t.get("checks", []):
        lines.append(f"  CONSTRAINT ck_{t['name']}_{ck['name']} CHECK ({ck['expr']})")
    w(f"CREATE TABLE {TS(t)}.{t['name']} (\n" + ",\n".join(lines) + "\n);")
    if t.get("seed"):
        cols = t.get("seed_cols") or [c["name"] for c in t["columns"][:3]]
        for row in t["seed"]:
            w(f"INSERT INTO {TS(t)}.{t['name']} ({', '.join(cols)}) VALUES ({', '.join('NULL' if v is None else (q(v) if isinstance(v,str) else str(v)) for v in row)});")
    w("")

w("-- ===== 외래키 =====")
for t in M["tables"]:
    for c in t["columns"]:
        if c.get("fk"):
            f = c["fk"]
            w(f"ALTER TABLE {TS(t)}.{t['name']} ADD CONSTRAINT fk_{t['name']}_{f['table']}_{c['name']} FOREIGN KEY ({c['name']}) REFERENCES {TSN(f['table'])}.{f['table']} ({f['column']}) ON DELETE {f['on_delete']} ON UPDATE NO ACTION;")
w("\n-- ===== FK 인덱스 =====")
for t in M["tables"]:
    pk = pkcols(t)
    for c in t["columns"]:
        if c.get("fk") and not (c["name"] == pk[0] if len(pk) == 1 else False):
            w(f"CREATE INDEX ix_{t['name']}_{c['name']} ON {TS(t)}.{t['name']} ({c['name']});")
w("\n-- ===== COMMENT (한글 논리명) =====")
for t in M["tables"]:
    w(f"COMMENT ON TABLE {TS(t)}.{t['name']} IS {q(t['kor'] + ' — ' + t['description'])};")
    for c in t["columns"]:
        d = c["kor"] + (" — " + c["description"] if c.get("description") else "") + (" [개인정보]" if c.get("pi") else "") \
            + (f" [암호화:{CRYPTO_LABEL[c['crypto']]}]" if c.get("crypto") else "")
        w(f"COMMENT ON COLUMN {TS(t)}.{t['name']}.{c['name']} IS {q(d)};")
if M.get("access_policies"):
    w("\n-- ===== 접근 정책 (data-platform 설치 시 meta.object_policy 에 등록) =====")
    w(f"DO $$ BEGIN\n  IF to_regclass('{MS}.object_policy') IS NOT NULL THEN")
    for p in M["access_policies"]:
        g = q(p["min_grade"]) if p.get("min_grade") else "NULL"; cg = q(p["category"]) if p.get("category") else "NULL"
        w(f"    INSERT INTO {MS}.object_policy (schema_nm, table_nm, min_grade_cd, ctgr_cd) VALUES ({q(p['schema'])}, {q(p.get('table') or '*')}, {g}, {cg})"
          f" ON CONFLICT (schema_nm, table_nm) DO UPDATE SET min_grade_cd = EXCLUDED.min_grade_cd, ctgr_cd = EXCLUDED.ctgr_cd, updated_at = now();")
    w("  END IF;\nEND $$;")
w("\nCOMMIT;")
open(OUT("02_ddl_postgresql.sql"), "w", encoding="utf-8").write("\n".join(L))

# ---------------- 2. pgModeler .dbm ----------------
def dbm_type(c):
    tt = ctype(c)
    m = re.match(r"(\w+)\((\d+)\)", tt)
    if m: return f'<type name="{m.group(1)}" length="{m.group(2)}"/>'
    return f'<type name="{tt}" length="0"/>'

X = []
x = X.append
n = len(M["tables"])
x('<?xml version="1.0" encoding="UTF-8"?>')
x('<!--\nCAUTION: Do not modify this file unless you know what you are doing.\n Unexpected results may occur if the code is changed deliberately.\n-->')
x(f'<dbmodel pgmodeler-ver="{A.pgmodeler_ver}" use-changelog="false" max-obj-count="{n+8}" last-position="0,0" last-zoom="1" scene-rect="0,0,2600,1600" default-schema="public" default-owner="postgres" layers="Default layer" active-layers="0" layer-name-colors="#000000" layer-rect-colors="#b4b4b4" show-layer-names="false" show-layer-rects="false">')
x(f'<database name="{S}_db" is-template="false" allow-conns="true">\n</database>')
x('<schema name="public" layers="0" fill-color="#e1e1e1" name-color="#000000" sql-disabled="true">\n</schema>')
x(f'<schema name="{S}" layers="0" rect-visible="true" fill-color="#e1e1e1" name-color="#000000">\n\t<role name="postgres"/>\n</schema>')
for _s in EXTRA:
    x(f'<schema name="{_s}" layers="0" rect-visible="true" fill-color="{"#fbe0e5" if _s == PIS else "#e1e1e1"}" name-color="#000000">\n\t<role name="postgres"/>\n</schema>')

# 배치: 기준(위) / 거래(중간) / PI(아래), 코드는 오른쪽
pos = {}
rows = {"코드": 60, "기준": 420, "거래": 820, "이력": 820, "PI": 1200}
colx = {"코드": 2000, "기준": 80, "거래": 300, "이력": 1500, "PI": 80}
for t in M["tables"]:
    k = t["kind"]; pos[t["name"]] = (colx[k], rows[k]); colx[k] += 440

for t in M["tables"]:
    ts = TS(t)
    x(f'<table name="{t["name"]}" alias="{escape(t["kor"])}" layers="0" collapse-mode="2" max-obj-count="{len(t["columns"])+4}" z-value="0">')
    x(f'\t<schema name="{ts}"/>\n\t<role name="postgres"/>')
    x(f'\t<comment> <![CDATA[{t["kor"]} — {t["description"]}]]> </comment>')
    px, py = pos[t["name"]]
    x(f'\t<position x="{px}" y="{py}"/>')
    for c in t["columns"]:
        attrs = f'name="{c["name"]}" alias="{escape(c["kor"])}"' + (' not-null="true"' if not c["nullable"] else "") + (' identity-type="ALWAYS"' if c.get("identity") else "")
        if c.get("default") is not None and not c.get("identity"):
            attrs += f' default-value="{escape(str(c["default"]))}"'
        x(f'\t<column {attrs}>\n\t\t{dbm_type(c)}\n\t\t<comment> <![CDATA[{c["kor"]}]]> </comment>\n\t</column>')
    x(f'\t<constraint name="pk_{t["name"]}" type="pk-constr" table="{ts}.{t["name"]}">\n\t\t<columns names="{",".join(pkcols(t))}" ref-type="src-columns"/>\n\t</constraint>')
    for c in t["columns"]:
        if c.get("unique"):
            x(f'\t<constraint name="uk_{t["name"]}_{c["name"]}" type="uq-constr" table="{ts}.{t["name"]}">\n\t\t<columns names="{c["name"]}" ref-type="src-columns"/>\n\t</constraint>')
        chk = DOM[c["domain"]]["check"]
        if chk:
            x(f'\t<constraint name="ck_{t["name"]}_{c["name"]}" type="ck-constr" table="{ts}.{t["name"]}">\n\t\t<expression> <![CDATA[{c["name"]} {chk}]]> </expression>\n\t</constraint>')
    for u in t.get("uniques", []):
        x(f'\t<constraint name="uk_{t["name"]}_{"_".join(u)}" type="uq-constr" table="{ts}.{t["name"]}">\n\t\t<columns names="{",".join(u)}" ref-type="src-columns"/>\n\t</constraint>')
    for ck in t.get("checks", []):
        x(f'\t<constraint name="ck_{t["name"]}_{ck["name"]}" type="ck-constr" table="{ts}.{t["name"]}">\n\t\t<expression> <![CDATA[{ck["expr"]}]]> </expression>\n\t</constraint>')
    x('</table>')

for t in M["tables"]:
    for c in t["columns"]:
        if c.get("fk"):
            f = c["fk"]
            x(f'<constraint name="fk_{t["name"]}_{f["table"]}_{c["name"]}" type="fk-constr" comparison-type="MATCH SIMPLE" upd-action="NO ACTION" del-action="{f["on_delete"]}" ref-table="{TSN(f["table"])}.{f["table"]}" table="{TS(t)}.{t["name"]}">')
            x(f'\t<columns names="{c["name"]}" ref-type="src-columns"/>\n\t<columns names="{f["column"]}" ref-type="dst-columns"/>\n</constraint>')
for t in M["tables"]:
    for c in t["columns"]:
        if c.get("fk"):
            f = c["fk"]
            rel_alias = f'{t["kor"]}→{TBL[f["table"]]["kor"]}'
            x(f'<relationship name="rel_{t["name"]}_{f["table"]}_{c["name"]}" alias="{escape(rel_alias)}" type="relfk" layers="0" src-table="{TS(t)}.{t["name"]}" dst-table="{TSN(f["table"])}.{f["table"]}" src-required="false" dst-required="false"/>')
x('</dbmodel>')
open(OUT("03_model.dbm"), "w", encoding="utf-8").write("\n".join(X))

# ---------------- 3. Excel 정의서 ----------------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
HDR = PatternFill("solid", fgColor="E4E7F7"); HF = Font(name="Malgun Gothic", bold=True, color="3A3F86", size=10)
BF = Font(name="Malgun Gothic", size=10); PIF = PatternFill("solid", fgColor="FBE0E5")
TH = Side(style="thin", color="D9D6E6"); BD = Border(left=TH, right=TH, top=TH, bottom=TH)
GROUP = PatternFill("solid", fgColor="F5F4FB")

def sheet(name, header, rows, widths=None, pi_col=None, group_col=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.fill = HDR; c.font = HF; c.alignment = Alignment(vertical="center", wrap_text=True); c.border = BD
    prev = None
    for r in rows:
        ws.append(r)
        row = ws[ws.max_row]
        for c in row:
            c.font = BF; c.border = BD; c.alignment = Alignment(vertical="center", wrap_text=True)
        if pi_col is not None and r[pi_col] in ("Y", True):
            for c in row: c.fill = PIF
        if group_col is not None and r[group_col] != prev:
            prev = r[group_col]
            if ws.max_row > 2:
                for c in row: c.border = Border(left=TH, right=TH, top=Side(style="medium", color="8E93E6"), bottom=TH)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24
    for i, wd in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(i)].width = wd
    return ws

# 0 개요
ws = wb.active; ws.title = "개요"
n_cols = sum(len(t["columns"]) for t in M["tables"]); n_pi = sum(1 for t in M["tables"] for c in t["columns"] if c.get("pi"))
info = [("프로젝트", M["project"]), ("버전", M["version"]), ("생성일", M["generated_at"]), ("DB", f"{M['db']} {M['db_version']}"),
        ("스키마", S if not EXTRA else ", ".join(SCHEMAS)), ("사전 스키마", MS), ("모델링 툴", M["modeling_tool"]), ("테이블 수", len(M["tables"])), ("컬럼 수", n_cols),
        ("개인정보 컬럼 수", n_pi), ("관계(FK) 수", len(M["relationships"])), ("", ""), ("가정", "\n".join("· " + a for a in M["assumptions"])),
        ("열린 질문", "\n".join("· " + a for a in M["open_questions"]))]
if PIS: info.insert(6, ("개인정보 스키마", PIS))
ws.append(["항목", "내용"]);
for c in ws[1]: c.fill = HDR; c.font = HF; c.border = BD
for k, v in info:
    ws.append([k, v])
    for c in ws[ws.max_row]: c.font = BF; c.border = BD; c.alignment = Alignment(vertical="top", wrap_text=True)
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 90

# 1 테이블목록
sheet("테이블목록", ["스키마", "테이블명", "논리명", "구분", "설명", "PI여부", "컬럼수"],
      [[TS(t), t["name"], t["kor"], t["kind"], t["description"], "Y" if t["is_pi"] else "", len(t["columns"])] for t in M["tables"]],
      [8, 24, 18, 8, 60, 8, 8], pi_col=5)

# 2 컬럼명세
rows = []
for t in M["tables"]:
    for i, c in enumerate(t["columns"], 1):
        d = DOM[c["domain"]]
        fk = f"{c['fk']['table']}.{c['fk']['column']}" if c.get("fk") else ""
        rows.append([t["name"], t["kor"], i, c["name"], c["kor"], c["domain"], d["type"],
                     "N" if not c["nullable"] else "Y", "PK" if c.get("pk") else "", fk, "UK" if c.get("unique") else "",
                     ("IDENTITY" if c.get("identity") else c.get("default", "")) or "", d["check"] or "", "Y" if c.get("pi") else "",
                     c.get("description", "") or ""])
sheet("컬럼명세", ["테이블명", "테이블논리명", "순번", "컬럼명", "논리명", "도메인", "자료형", "NULL", "PK", "FK(참조)", "UNIQUE", "기본값", "CHECK", "PI", "설명"],
      rows, [22, 14, 6, 16, 14, 11, 13, 6, 5, 28, 8, 12, 34, 5, 36], pi_col=13, group_col=0)

# 3 관계(FK)
sheet("관계(FK)", ["제약명", "자식테이블", "자식컬럼", "부모테이블", "부모컬럼", "관계", "필수", "삭제규칙", "갱신규칙", "업무의미"],
      [[f"fk_{r['from']}_{r['to']}_{r['from_col']}", r["from"], r["from_col"], r["to"], r["to_col"], r["type"],
        "N" if r["optional"] else "Y", r["on_delete"], "NO ACTION", r["meaning"]] for r in M["relationships"]],
      [40, 22, 14, 22, 14, 6, 6, 10, 10, 40])

# 4 용어사전
ws = sheet("용어사전", ["[단어] 한글", "약어", "영문", "분류어", "설명"],
      [[x["kor"], x["abbr"], x["full"], "Y" if x["is_class"] else "", x["desc"]] for x in M["words"]], [14, 10, 22, 8, 40])
ws.append([]); ws.append(["[용어] 한글명", "물리명", "단어경로", "도메인", "설명"])
for c in ws[ws.max_row]: c.fill = HDR; c.font = HF; c.border = BD
for t in M["terms"]:
    ws.append([t["kor"], t["phys"], t["word_path"], t["domain"] or "", t.get("desc") or ""])
    for c in ws[ws.max_row]: c.font = BF; c.border = BD
ws.append([]); ws.append(["[도메인] 이름", "자료형", "검증규칙", "예시", "설명"])
for c in ws[ws.max_row]: c.fill = HDR; c.font = HF; c.border = BD
for d in M["domains"]:
    ws.append([d["name"], d["type"], d["check"] or "", d["example"], d["desc"]])
    for c in ws[ws.max_row]: c.font = BF; c.border = BD
ws.auto_filter.ref = None

# 5 개인정보
def _pi_now(t, c):
    if c.get("crypto"): return f"{TS(t)}.{t['name']} 1:1 분리 · {CRYPTO_LABEL[c['crypto']]}"
    return "pi_ 테이블 1:1 분리(평문)"
def _pi_sug(c):
    if c.get("crypto"): return "마스킹 뷰 · 보존기간 파기 · 등급/분류 접근정책(개인정보) · 키 교체 주기"
    return "pgcrypto 암호화 · 마스킹 뷰 · 별도 스키마/ROLE · 보존기간 파기 · RLS"
rows = [[t["name"], c["name"], c["kor"], c.get("pi_type", "개인정보"), "Y" if c.get("sensitive") else "",
         _pi_now(t, c), _pi_sug(c)]
        for t in M["tables"] for c in t["columns"] if c.get("pi")]
sheet("개인정보", ["테이블", "컬럼", "논리명", "개인정보 유형", "민감", "현재 처리", "제안"], rows, [16, 14, 12, 22, 6, 24, 60])

# 6 코드정의
rows = []
for t in M["tables"]:
    for row in t.get("seed", []):
        rows.append([t["name"], t["kor"], row[0], row[1], row[2]])
for d in M["domains"]:
    if d["check"] and d["check"].strip().upper().startswith("IN"):
        rows.append([f"(CHECK) {d['name']}", d["name"], d["check"], "고정값", ""])
sheet("코드정의", ["코드테이블", "논리명", "코드값", "코드명", "순서"], rows, [22, 18, 30, 14, 6])

# 7 검토메모
rows = [["DQ", n] for n in M["dq_notes"]] + [["PI 제안", n] for n in M["pi_suggestions"]] + \
       [["반정규화", n] for n in M.get("denormalization_notes", ["없음 — 3NF 유지"])] + \
       [["향후 변화", n] for n in M.get("change_scenarios", [])]
sheet("검토메모", ["구분", "내용"], rows, [12, 110])

# 8 접근정책 (선택)
if M.get("access_policies"):
    sheet("접근정책", ["스키마", "테이블", "최소등급", "분류", "의미"],
          [[p["schema"], p.get("table") or "*", p.get("min_grade") or "", p.get("category") or "", p.get("meaning", "")] for p in M["access_policies"]],
          [12, 24, 10, 12, 60])

wb.save(OUT("04_table_spec.xlsx"))
print("exported →", A.outdir, ": 02_ddl_postgresql.sql, 03_model.dbm, 04_table_spec.xlsx")
